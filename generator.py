"""
Answer generation module supporting both Gemini API and Local Hugging Face models.
Features:
- Dual backend: "gemini" (Google Gemini API via google-genai) and "local" (Qwen2.5-0.5B-Instruct / 1.5B / Flan-T5)
- Auto-fallback to local model with clear console message if Gemini API key is missing or invalid
- Token-by-token streaming into the Gradio chat interface for immediate responsiveness
- Heuristic follow-up rewriting (runs only on short queries or queries with pronouns)
- Reduced work per answer: top 3 chunks, ~400 char truncation, max 2 history turns, max 200 tokens
- Persistent disk caching for 5 example questions and Map-Reduce summaries
- Console timing logs for retrieval and generation benchmarks
"""

import os
import json
import logging
import re
import time
from threading import Thread, current_thread
from pathlib import Path
from typing import List, Dict, Any, Optional, Generator as PyGenerator

import numpy as np
import torch
from transformers import AutoTokenizer, AutoModelForSeq2SeqLM, AutoModelForCausalLM, AutoConfig, TextIteratorStreamer

from config import (
    BASE_DIR,
    GENERATOR_BACKEND,
    GEMINI_API_KEY,
    GEMINI_MODEL_NAME,
    LOCAL_MODEL_NAME,
    FALLBACK_MODEL_NAME,
    TOP_K,
    CHUNK_CONTEXT_CHAR_LIMIT,
    MAX_CHAT_HISTORY_TURNS,
    MAX_NEW_TOKENS,
    DO_SAMPLE,
    TEMPERATURE,
    SYSTEM_PROMPT,
    USER_PROMPT_TEMPLATE,
    MAP_SUMMARY_PROMPT,
    REDUCE_SUMMARY_DETAILED_PROMPT,
    REDUCE_SUMMARY_BRIEF_PROMPT,
    REWRITE_PROMPT_TEMPLATE,
    CHUNK_QUESTION_PROMPT,
    CACHE_DIR,
    SIMILARITY_THRESHOLD,
    PAGES_PER_SUMMARY_GROUP,
    CONFIDENCE_HIGH_THRESHOLD,
    CONFIDENCE_MED_THRESHOLD
)

try:
    import docx
    from docx.shared import Pt, RGBColor
except ImportError:
    docx = None


logger = logging.getLogger(__name__)


class Generator:
    """
    Manages LLM inference for grounded document Q&A, follow-up rewriting,
    validated question generation, and map-reduce summarization across Gemini & Local backends.
    """
    _instance: Optional["Generator"] = None

    def __init__(self, backend: str = GENERATOR_BACKEND, local_model_name: str = LOCAL_MODEL_NAME):
        self.requested_backend = backend
        self.backend = "local"
        self.local_model_name = local_model_name
        self.gemini_client = None
        self.tokenizer = None
        self.model = None
        self.is_seq2seq = False
        self._is_stopped = False
        self._summary_threads: Dict[str, Thread] = {}

        # 1. Attempt Gemini backend if requested and API key is present
        if backend == "gemini" and GEMINI_API_KEY:
            try:
                from google import genai
                client = genai.Client(api_key=GEMINI_API_KEY)
                # Quick verification call
                test_resp = client.models.generate_content(
                    model=GEMINI_MODEL_NAME,
                    contents="Hi"
                )
                if test_resp and test_resp.text:
                    self.gemini_client = client
                    self.backend = "gemini"
                    print(f"[GENERATOR] Active backend: GEMINI ({GEMINI_MODEL_NAME})")
            except Exception as e:
                print(f"[INFO] Gemini API unavailable or project access denied ({e}). Falling back to local backend: {local_model_name}")
                self.backend = "local"
        elif backend == "gemini" and not GEMINI_API_KEY:
            print(f"[INFO] No GEMINI_API_KEY found in .env. Falling back to local backend: {local_model_name}")
            self.backend = "local"

        # 2. Initialize local Hugging Face model if active backend is local
        if self.backend == "local":
            self._init_local_model(local_model_name)
            print(f"[GENERATOR] Active backend: LOCAL ({self.local_model_name})")

    def stop_generation(self) -> None:
        """Stops the active streaming generation thread."""
        self._is_stopped = True

    def _init_local_model(self, model_name: str) -> None:
        """Loads local tokenizer and model into CPU memory once at startup."""
        self.local_model_name = model_name
        try:
            logger.info("Loading local model: %s...", model_name)
            t0 = time.time()
            self.tokenizer = AutoTokenizer.from_pretrained(model_name)
            model_config = AutoConfig.from_pretrained(model_name)
            self.is_seq2seq = model_config.is_encoder_decoder

            if self.is_seq2seq:
                self.model = AutoModelForSeq2SeqLM.from_pretrained(
                    model_name,
                    dtype=torch.float32
                )
            else:
                self.model = AutoModelForCausalLM.from_pretrained(
                    model_name,
                    dtype="auto",
                    low_cpu_mem_usage=True
                )

            self.model.to("cpu")
            self.model.eval()
            logger.info("Local model loaded in %.2fs (Seq2Seq: %s)", time.time() - t0, self.is_seq2seq)
        except Exception as e:
            if model_name != FALLBACK_MODEL_NAME:
                logger.warning("Failed to load %s (%s). Falling back to %s...", model_name, e, FALLBACK_MODEL_NAME)
                self._init_local_model(FALLBACK_MODEL_NAME)
            else:
                raise e

    @classmethod
    def get_instance(cls, backend: str = GENERATOR_BACKEND, local_model_name: str = LOCAL_MODEL_NAME) -> "Generator":
        """Singleton accessor."""
        if cls._instance is None:
            cls._instance = cls(backend=backend, local_model_name=local_model_name)
        return cls._instance

    def format_context(self, retrieved_chunks: List[Dict[str, Any]]) -> str:
        """
        Formats context chunks with page citations.
        Work reduction: takes at most top 3 chunks and truncates each to ~400 characters.
        """
        if not retrieved_chunks:
            return ""
        formatted = []
        for chunk in retrieved_chunks[:TOP_K]:
            page_no = chunk.get("page_number", "?")
            text = chunk.get("text", "").strip()
            if len(text) > CHUNK_CONTEXT_CHAR_LIMIT:
                text = text[:CHUNK_CONTEXT_CHAR_LIMIT] + "..."
            formatted.append(f"[Page {page_no}]: {text}")
        return "\n\n".join(formatted)

    def format_page_citation(self, chunks: List[Dict[str, Any]]) -> str:
        """Generates sorted unique page numbers: 'Sources: pages X, Y'."""
        if not chunks:
            return ""
        pages = sorted(list(set(c["page_number"] for c in chunks[:TOP_K] if "page_number" in c)))
        if not pages:
            return ""
        if len(pages) == 1:
            return f"Sources: page {pages[0]}"
        return f"Sources: pages {', '.join(str(p) for p in pages)}"

    def should_rewrite_question(self, question: str) -> bool:
        """
        Speed optimization: only rewrite follow-up questions when the question is under 6 words
        or contains pronouns like it/that/this/they/those. Otherwise skips the model call.
        """
        clean = (question or "").strip()
        words = clean.split()
        if len(words) < 6:
            return True
        pronouns = {"it", "that", "this", "they", "those", "them", "their", "these", "its", "he", "she", "his", "her"}
        tokens = set(re.findall(r"\b\w+\b", clean.lower()))
        return bool(tokens & pronouns)

    def rewrite_question_with_history(
        self,
        question: str,
        history: List[Dict[str, str]]
    ) -> str:
        """
        Rewrites conversational follow-up questions into standalone questions
        using the last 2 chat turns. Skips rewriting when unnecessary.
        """
        clean_question = (question or "").strip()
        if not clean_question or not history:
            return clean_question

        # Speed optimization check
        if not self.should_rewrite_question(clean_question):
            return clean_question

        # Limit to at most the last 2 chat turns (up to 4 messages)
        recent_turns = history[-(MAX_CHAT_HISTORY_TURNS * 2):] if len(history) > (MAX_CHAT_HISTORY_TURNS * 2) else history
        history_lines = []
        for msg in recent_turns:
            role = "User" if msg.get("role") == "user" else "Assistant"
            content = msg.get("content", "").split("\n\nSources:")[0].strip()
            if len(content) > 180:
                content = content[:180] + "..."
            history_lines.append(f"{role}: {content}")

        history_str = "\n".join(history_lines)
        prompt = REWRITE_PROMPT_TEMPLATE.format(history=history_str, question=clean_question)

        try:
            rewritten = self._run_inference(prompt, max_tokens=64)
            rewritten_clean = re.sub(r'^(Standalone Question:\s*|\"|\')', '', rewritten).strip(' "\'')
            if rewritten_clean and len(rewritten_clean) > 5:
                logger.info("Rewrote question '%s' -> '%s'", clean_question, rewritten_clean)
                return rewritten_clean
        except Exception as e:
            logger.warning("Error rewriting question: %s. Using original.", e)

        return clean_question

    def _run_inference(self, prompt: str, max_tokens: int = MAX_NEW_TOKENS) -> str:
        """Synchronous inference runner supporting Gemini or Local model."""
        if self.backend == "gemini" and self.gemini_client:
            try:
                resp = self.gemini_client.models.generate_content(
                    model=GEMINI_MODEL_NAME,
                    contents=prompt
                )
                return (resp.text or "").strip()
            except Exception as e:
                logger.warning("Gemini call failed (%s). Falling back to local inference.", e)

        # Local HF inference
        inputs = self.tokenizer(prompt, return_tensors="pt", truncation=True, max_length=1536).to("cpu")
        with torch.no_grad():
            gen_kwargs = {
                "max_new_tokens": max_tokens,
                "do_sample": DO_SAMPLE,
            }
            if DO_SAMPLE and TEMPERATURE > 0.0:
                gen_kwargs["temperature"] = TEMPERATURE

            output_tokens = self.model.generate(
                input_ids=inputs["input_ids"],
                attention_mask=inputs.get("attention_mask"),
                **gen_kwargs
            )
            if self.is_seq2seq:
                raw_text = self.tokenizer.decode(output_tokens[0], skip_special_tokens=True)
            else:
                generated = output_tokens[0][inputs["input_ids"].shape[-1]:]
                raw_text = self.tokenizer.decode(generated, skip_special_tokens=True)

        return raw_text.strip()

    def generate_answer_stream(
        self,
        question: str,
        retrieved_chunks: List[Dict[str, Any]],
        answer_style: str = "Detailed (3-6 sentences or bullets)"
    ) -> PyGenerator[str, None, None]:
        """
        Streams generated answers token by token into the chat interface for immediate user feedback.
        Supports 'Short' (1-2 sentences, max 80 tokens) and 'Detailed' (3-6 sentences, max 200 tokens) styles.
        Appends canonical page citations at the end.
        """
        clean_question = (question or "").strip()
        if not clean_question:
            yield "Please enter a valid question."
            return

        if not retrieved_chunks:
            yield "Answer not found in the document."
            return

        self._is_stopped = False
        is_short = "short" in (answer_style or "").lower()
        tokens_limit = 80 if is_short else MAX_NEW_TOKENS

        if is_short:
            sys_prompt = (
                "You are a helpful assistant for question answering over PDF documents.\n"
                "Answer ONLY from the provided context below.\n"
                "Provide your answer in 1-2 concise, factual sentences. Be direct and avoid introductory phrases.\n"
                "Do not invent facts, extrapolate, or assume outside information.\n"
                "If the context does not contain the answer, reply exactly: 'Answer not found in the document.'"
            )
        else:
            sys_prompt = SYSTEM_PROMPT

        context_str = self.format_context(retrieved_chunks)

        # Build prompt
        if self.backend == "local" and hasattr(self.tokenizer, "apply_chat_template") and self.tokenizer.chat_template and not self.is_seq2seq:
            messages = [
                {"role": "system", "content": sys_prompt},
                {"role": "user", "content": USER_PROMPT_TEMPLATE.format(context=context_str, question=clean_question)}
            ]
            prompt = self.tokenizer.apply_chat_template(messages, tokenize=False, add_generation_prompt=True)
        else:
            prompt = (
                f"{sys_prompt}\n\n"
                f"Context:\n{context_str}\n\n"
                f"Question: {clean_question}\n\n"
                "Answer:"
            )

        accumulated_text = ""

        # --- STREAMING: Gemini Backend ---
        if self.backend == "gemini" and self.gemini_client:
            try:
                stream = self.gemini_client.models.generate_content_stream(
                    model=GEMINI_MODEL_NAME,
                    contents=prompt
                )
                for chunk in stream:
                    if self._is_stopped:
                        break
                    if chunk.text:
                        accumulated_text += chunk.text
                        yield accumulated_text
            except Exception as e:
                logger.warning("Gemini streaming failed (%s). Falling back to local.", e)
                accumulated_text = ""

        # --- STREAMING: Local HF Backend ---
        if not accumulated_text:
            inputs = self.tokenizer(prompt, return_tensors="pt", truncation=True, max_length=1536).to("cpu")
            streamer = TextIteratorStreamer(self.tokenizer, skip_prompt=True, skip_special_tokens=True)

            gen_kwargs = dict(
                input_ids=inputs["input_ids"],
                attention_mask=inputs.get("attention_mask"),
                max_new_tokens=tokens_limit,
                do_sample=DO_SAMPLE,
                streamer=streamer
            )

            thread = Thread(target=self.model.generate, kwargs=gen_kwargs)
            thread.daemon = True
            thread.start()

            for new_text in streamer:
                if self._is_stopped:
                    break
                accumulated_text += new_text
                yield accumulated_text

        # Check refusal and format citation & confidence
        norm = accumulated_text.lower().strip()
        if not accumulated_text or "not found in the document" in norm or "answer not found" in norm:
            yield "Answer not found in the document."
            return

        clean_answer = re.sub(r'Sources?:.*$', '', accumulated_text, flags=re.IGNORECASE).strip()
        if self._is_stopped:
            clean_answer += " *(generation stopped)*"

        citation = self.format_page_citation(retrieved_chunks)
        confidence = self.compute_confidence(retrieved_chunks)

        meta_parts = []
        if citation:
            meta_parts.append(citation)
        if confidence:
            meta_parts.append(confidence)

        if meta_parts:
            meta_str = " • ".join(meta_parts)
            yield f"{clean_answer}\n\n<small style='color: #6B7280; font-size: 13.5px;'>{meta_str}</small>"
        else:
            yield clean_answer

    def compute_confidence(self, retrieved_chunks: List[Dict[str, Any]]) -> str:
        """
        Computes confidence level based on top retrieval semantic similarity.
        High (>=0.45), Medium (>=0.30), Low (<0.30).
        """
        if not retrieved_chunks:
            return ""
        top_sim = retrieved_chunks[0].get("similarity_score", 0.0)
        if top_sim >= CONFIDENCE_HIGH_THRESHOLD:
            return "Confidence: High"
        elif top_sim >= CONFIDENCE_MED_THRESHOLD:
            return "Confidence: Medium"
        else:
            return "Confidence: Low (Please verify on the cited page.)"

    def generate_answer(
        self,
        question: str,
        retrieved_chunks: List[Dict[str, Any]]
    ) -> str:
        """Synchronous wrapper for evaluation benchmarks."""
        final_answer = ""
        for partial in self.generate_answer_stream(question, retrieved_chunks):
            final_answer = partial
        # Clean HTML tag from citation for evaluate.py text matching
        return re.sub(r"<[^>]+>", "", final_answer).strip()


    def generate_5_example_questions(
        self,
        retriever: Any,
        embedder: Any,
        force_regenerate: bool = False
    ) -> List[str]:
        """
        Generates 5 verified example questions across the document.
        Uses SHA-256 disk caching for instant repeated PDF loads.
        """
        if not retriever or not retriever.chunks:
            return []

        pdf_hash = getattr(retriever, "pdf_hash", None)
        cache_path = (CACHE_DIR / f"{pdf_hash}_example_questions.json") if pdf_hash else None

        if not force_regenerate and cache_path and cache_path.exists():
            try:
                with open(cache_path, "r", encoding="utf-8") as f:
                    cached = json.load(f)
                if isinstance(cached, list) and len(cached) == 5:
                    logger.info("Loaded 5 example questions from disk cache.")
                    return cached
            except Exception as e:
                logger.warning("Could not read example questions cache: %s", e)

        # 1. Pick 8 representative chunks spread evenly
        chunks = retriever.get_representative_chunks(count=8)
        if not chunks:
            return []

        # 2. Ask LLM to write one question per chunk
        excerpts_text = []
        for i, c in enumerate(chunks, start=1):
            clean_chunk = c.get("text", "").strip().replace("\n", " ")
            if len(clean_chunk) > 250:
                clean_chunk = clean_chunk[:250].rsplit(" ", 1)[0]
            excerpts_text.append(f"Passage {i} (p. {c.get('page_number', '?')}):\n\"{clean_chunk}\"")

        prompt = (
            "For each passage below, write ONE clear, specific, concise question directly answered by it.\n"
            "Rules:\n"
            "- Start each question with 'What', 'How', 'Why', 'Define', or 'Compare'.\n"
            "- Each question must be complete, directly answered by the passage, and at most 15 words.\n"
            "- Each question must be a complete sentence ending with '?'. Do NOT use ellipsis (...) and do NOT cut off questions.\n"
            "- Do NOT ask yes/no or vague questions.\n"
            "- Output ONLY numbered lines: '1. Question', '2. Question', etc.\n\n"
            f"{chr(10).join(excerpts_text)}\n\n"
            "Questions:"
        )

        candidates = []
        try:
            raw_output = self._run_inference(prompt, max_tokens=400)
            for line in raw_output.split("\n"):
                line_clean = line.strip()
                q_text = re.sub(r'^\d+[\.\)]\s*', '', line_clean).strip(' "\'')
                if (
                    12 <= len(q_text) <= 120
                    and q_text.endswith("?")
                    and not q_text.endswith("...")
                    and not q_text.endswith("..")
                    and "..." not in q_text
                    and len(q_text.split()) <= 16
                ):
                    q_lower = q_text.lower()
                    if not any(bad in q_lower for bad in ["what is this document", "what does this document discuss", "is there"]):
                        candidates.append(q_text)
        except Exception as e:
            logger.warning("Error generating candidate questions: %s", e)

        # 3. Validate each question through retriever
        valid_candidates = []
        for q in candidates:
            try:
                ret_chunks = retriever.retrieve(q, top_k=1, similarity_threshold=SIMILARITY_THRESHOLD)
                if ret_chunks and ret_chunks[0]["similarity_score"] >= SIMILARITY_THRESHOLD:
                    valid_candidates.append((q, ret_chunks[0]["similarity_score"]))
            except Exception:
                continue

        valid_candidates.sort(key=lambda x: x[1], reverse=True)

        # 4. Deduplicate (embedding cosine similarity > 0.85)
        selected = []
        if valid_candidates:
            q_texts = [q for q, _ in valid_candidates]
            embeddings = embedder.embed_texts(q_texts)
            sim_matrix = np.dot(embeddings, embeddings.T)

            selected_indices = []
            for i, (q, _) in enumerate(valid_candidates):
                is_duplicate = False
                for prev_i in selected_indices:
                    if sim_matrix[i, prev_i] > 0.85:
                        is_duplicate = True
                        break
                if not is_duplicate:
                    selected.append(q)
                    selected_indices.append(i)
                if len(selected) == 5:
                    break

        # 5. Fallback to headings if needed
        if len(selected) < 5:
            headings = retriever.get_document_headings()
            fallback_templates = [
                "What are the main concepts of {h}?",
                "How does the document describe {h}?",
                "What key details are provided about {h}?",
                "Why is {h} significant in this document?",
                "Define the role of {h}."
            ]
            t_idx = 0
            for h in headings:
                if len(selected) >= 5:
                    break
                h_clean = h.strip()
                if len(h_clean) > 40:
                    h_clean = h_clean[:40].rsplit(" ", 1)[0]
                if len(h_clean) >= 4 and not any(h_clean.lower() in s.lower() for s in selected):
                    template = fallback_templates[t_idx % len(fallback_templates)]
                    t_idx += 1
                    fallback_q = template.format(h=h_clean)
                    selected.append(fallback_q)

        # Final generic fallbacks
        generic_fallbacks = [
            "What is the main objective discussed in the document?",
            "What are the primary technical findings presented?",
            "What methodologies or approaches were evaluated?",
            "What key quantitative metrics are reported?",
            "What conclusions or future directions are outlined?"
        ]
        for gf in generic_fallbacks:
            if len(selected) >= 5:
                break
            if gf not in selected:
                selected.append(gf)

        selected = selected[:5]

        # Cache results
        if cache_path:
            try:
                with open(cache_path, "w", encoding="utf-8") as f:
                    json.dump(selected, f, ensure_ascii=False, indent=2)
            except Exception as e:
                logger.warning("Could not cache example questions: %s", e)

        return selected

    def precompute_summary_async(
        self,
        retriever: Any,
        length_option: str = "Detailed (~200 words)"
    ) -> None:
        """
        Launches a background daemon thread to precompute the Detailed summary right after upload.
        Caches the result to disk so that clicking 'Summarize' is instant.
        """
        if not retriever or not retriever.chunks:
            return
        pdf_hash = getattr(retriever, "pdf_hash", None)
        if not pdf_hash:
            return

        is_brief = "brief" in length_option.lower()
        length_key = "brief" if is_brief else "detailed"
        cache_path = CACHE_DIR / f"{pdf_hash}_summary_{length_key}.json"
        if cache_path.exists():
            return

        def _bg_worker():
            logger.info("Starting background summary precomputation for %s...", pdf_hash[:8])
            t0 = time.time()
            try:
                self.generate_map_reduce_summary(retriever, length_option=length_option)
                logger.info("Background summary precomputation completed in %.2fs", time.time() - t0)
            except Exception as e:
                logger.warning("Background summary precomputation failed: %s", e)

        th = Thread(target=_bg_worker, daemon=True)
        self._summary_threads[pdf_hash] = th
        th.start()

    def generate_summary_stream(
        self,
        retriever: Any,
        length_option: str = "Detailed (~200 words)",
        force_regenerate: bool = False
    ) -> PyGenerator[str, None, None]:
        """
        Streams document summary token-by-token:
        1. Instant return if cached on disk.
        2. Waits and yields 'Preparing summary...' if background thread is active.
        3. For PDFs <= 15 pages, uses fast single-call summarization with evenly sampled chunks.
        4. For PDFs > 15 pages, uses Map-Reduce summarization.
        5. Caps max_new_tokens (250 for Detailed, 120 for Brief).
        """
        if not retriever or not retriever.chunks:
            yield "Please upload a valid PDF document to summarize."
            return

        is_brief = "brief" in length_option.lower()
        length_key = "brief" if is_brief else "detailed"
        pdf_hash = getattr(retriever, "pdf_hash", None)
        cache_path = (CACHE_DIR / f"{pdf_hash}_summary_{length_key}.json") if pdf_hash else None

        # 1. Instant cache check
        if not force_regenerate and cache_path and cache_path.exists():
            try:
                with open(cache_path, "r", encoding="utf-8") as f:
                    cached = json.load(f)
                if isinstance(cached, str) and len(cached) > 20:
                    logger.info("Loaded %s summary from disk cache (instant).", length_key)
                    yield cached
                    return
            except Exception:
                pass

        # 2. Check if background precomputation is in progress
        if pdf_hash and pdf_hash in self._summary_threads:
            th = self._summary_threads[pdf_hash]
            if th != current_thread() and th.is_alive():
                yield "Preparing summary in background..."
                while th.is_alive():
                    time.sleep(0.4)
                if cache_path and cache_path.exists():
                    try:
                        with open(cache_path, "r", encoding="utf-8") as f:
                            cached = json.load(f)
                        if isinstance(cached, str) and len(cached) > 20:
                            yield cached
                            return
                    except Exception:
                        pass

        # 3. Choose summarization strategy
        total_pages = len(set(c.get("page_number", 1) for c in retriever.chunks))
        self._is_stopped = False
        max_tokens = 120 if is_brief else 250

        if total_pages <= 15:
            # Single-call fast summarization
            sample_chunks = retriever.get_representative_chunks(count=min(7, len(retriever.chunks)))
            snippets = []
            for c in sample_chunks:
                p = c.get("page_number", "?")
                txt = c.get("text", "").replace("\n", " ").strip()
                if len(txt) > 260:
                    txt = txt[:260] + "..."
                snippets.append(f"[p. {p}]: {txt}")
            combined_context = "\n\n".join(snippets)
            reduce_template = REDUCE_SUMMARY_BRIEF_PROMPT if is_brief else REDUCE_SUMMARY_DETAILED_PROMPT
            prompt = reduce_template.format(context=combined_context)
        else:
            # Map-Reduce for long PDFs (> 15 pages)
            groups = retriever.get_chunks_grouped_by_page_range(pages_per_group=PAGES_PER_SUMMARY_GROUP)
            map_summaries = []
            for g in groups:
                p_range = g["page_range"]
                ctx = g["text"][:900]
                p_map = MAP_SUMMARY_PROMPT.format(page_range=p_range, context=ctx)
                s_text = self._run_inference(p_map, max_tokens=100)
                s_clean = re.sub(r'^(Group Summary:\s*|\"|\')', '', s_text).strip(' "\'')
                map_summaries.append(f"[{p_range}]: {s_clean}")
            combined_summaries = "\n\n".join(map_summaries)
            reduce_template = REDUCE_SUMMARY_BRIEF_PROMPT if is_brief else REDUCE_SUMMARY_DETAILED_PROMPT
            prompt = reduce_template.format(context=combined_summaries)

        # 4. Stream generation
        accumulated = ""
        if self.backend == "local":
            inputs = self.tokenizer(prompt, return_tensors="pt", truncation=True, max_length=1536).to("cpu")
            streamer = TextIteratorStreamer(self.tokenizer, skip_prompt=True, skip_special_tokens=True)
            gen_kwargs = dict(
                input_ids=inputs["input_ids"],
                attention_mask=inputs.get("attention_mask"),
                max_new_tokens=max_tokens,
                do_sample=DO_SAMPLE,
                streamer=streamer
            )
            th = Thread(target=self.model.generate, kwargs=gen_kwargs)
            th.daemon = True
            th.start()
            for new_text in streamer:
                if self._is_stopped:
                    break
                accumulated += new_text
                clean_disp = re.sub(r'^(Structured Summary:\s*|Brief Summary:\s*)', '', accumulated).strip()
                yield clean_disp
        else:
            try:
                stream = self.gemini_client.models.generate_content_stream(model=GEMINI_MODEL_NAME, contents=prompt)
                for chunk in stream:
                    if self._is_stopped:
                        break
                    if chunk.text:
                        accumulated += chunk.text
                        clean_disp = re.sub(r'^(Structured Summary:\s*|Brief Summary:\s*)', '', accumulated).strip()
                        yield clean_disp
            except Exception:
                pass

        final_summary = re.sub(r'^(Structured Summary:\s*|Brief Summary:\s*)', '', accumulated).strip()
        if cache_path and final_summary and len(final_summary) > 20:
            try:
                with open(cache_path, "w", encoding="utf-8") as f:
                    json.dump(final_summary, f, ensure_ascii=False, indent=2)
            except Exception as e:
                logger.warning("Could not cache summary: %s", e)

    def generate_map_reduce_summary(
        self,
        retriever: Any,
        length_option: str = "Detailed (~200 words)",
        force_regenerate: bool = False
    ) -> str:
        """Synchronous wrapper for summarization benchmark and background precomputation."""
        final_summary = ""
        for partial in self.generate_summary_stream(retriever, length_option=length_option, force_regenerate=force_regenerate):
            final_summary = partial
        return final_summary


def _safe_extract_text(content: Any) -> str:
    """Safely extracts a flat text string from str, list of chunks/dicts, or dict."""
    if isinstance(content, str):
        return content
    if isinstance(content, (list, tuple)):
        parts = []
        for item in content:
            if isinstance(item, dict):
                parts.append(str(item.get("text", item.get("content", str(item)))))
            else:
                parts.append(str(item))
        return " ".join(parts)
    if isinstance(content, dict):
        return str(content.get("text", content.get("content", str(content))))
    return str(content or "")


def export_chat_transcript(
    history: List[Dict[str, Any]],
    doc_name: str = "Document",
    summary_text: Optional[str] = None,
    format_type: str = "docx"
) -> Optional[str]:
    """
    Exports conversation history into a nicely formatted .docx or .txt document in a temp folder.
    Includes document name, date, summary (if available), each Q&A turn, and citations.
    """
    if not history:
        return None

    import tempfile
    temp_dir = Path(tempfile.gettempdir()) / "nlp_rag_exports"
    temp_dir.mkdir(parents=True, exist_ok=True)

    clean_doc = Path(doc_name).stem if doc_name else "document"
    clean_doc = re.sub(r"[^\w\-]", "_", clean_doc)
    timestamp_str = time.strftime("%Y-%m-%d %H:%M:%S")
    file_timestamp = time.strftime("%Y%m%d_%H%M%S")

    is_docx = "docx" in format_type.lower() and docx is not None
    out_filename = f"chat_{clean_doc}_{file_timestamp}.{'docx' if is_docx else 'txt'}"
    out_path = temp_dir / out_filename

    if is_docx:
        try:
            doc = docx.Document()
            title = doc.add_heading("PDF Question Answering Transcript", level=0)
            if title.runs:
                title.runs[0].font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

            meta = doc.add_paragraph()
            meta.add_run("Document: ").bold = True
            meta.add_run(f"{doc_name}\n")
            meta.add_run("Exported: ").bold = True
            meta.add_run(f"{timestamp_str}\n")

            if summary_text and summary_text.strip():
                clean_sum = _safe_extract_text(summary_text)
                doc.add_heading("Executive Summary", level=2)
                p_sum = doc.add_paragraph(clean_sum.strip())
                if p_sum.runs:
                    p_sum.runs[0].font.size = Pt(10.5)

            doc.add_heading("Conversation", level=2)

            for msg in history:
                role = msg.get("role", "")
                raw_content = _safe_extract_text(msg.get("content", ""))
                clean_content = re.sub(r"<[^>]+>", "", raw_content).strip()

                if role == "user":
                    p_q = doc.add_paragraph()
                    r_q = p_q.add_run(f"Q: {clean_content}")
                    r_q.bold = True
                    r_q.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
                elif role == "assistant":
                    p_a = doc.add_paragraph()
                    r_a = p_a.add_run(f"A: {clean_content}")
                    r_a.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
                    p_a.paragraph_format.space_after = Pt(8)

            doc.save(str(out_path))
            return str(out_path)
        except Exception as e:
            logger.warning("DOCX export failed (%s), falling back to TXT", e)
            out_path = temp_dir / f"chat_{clean_doc}_{file_timestamp}.txt"

    # TXT Export fallback
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("=" * 60 + "\n")
        f.write("PDF QUESTION ANSWERING TRANSCRIPT\n")
        f.write(f"Document: {doc_name}\n")
        f.write(f"Date:     {timestamp_str}\n")
        f.write("=" * 60 + "\n\n")

        if summary_text and summary_text.strip():
            f.write("--- DOCUMENT SUMMARY ---\n")
            f.write(_safe_extract_text(summary_text).strip() + "\n\n")

        f.write("--- QUESTIONS & ANSWERS ---\n\n")
        for msg in history:
            role = "USER" if msg.get("role") == "user" else "ASSISTANT"
            clean_content = re.sub(r"<[^>]+>", "", _safe_extract_text(msg.get("content", ""))).strip()
            f.write(f"[{role}]:\n{clean_content}\n\n")

    return str(out_path)


def log_feedback(
    doc_name: str,
    question: str,
    answer: str,
    pages: str,
    rating: str
) -> None:
    """
    Appends user feedback (like/dislike) to feedback_log.csv.
    """
    import csv
    csv_path = BASE_DIR / "feedback_log.csv"
    file_exists = csv_path.exists()

    clean_q = question.replace("\n", " ").strip()
    clean_a = re.sub(r"<[^>]+>", "", answer).replace("\n", " ").strip()

    with open(csv_path, "a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if not file_exists:
            writer.writerow(["Timestamp", "Document", "Question", "Answer", "Pages", "Rating"])
        writer.writerow([
            time.strftime("%Y-%m-%d %H:%M:%S"),
            doc_name,
            clean_q,
            clean_a,
            pages,
            rating
        ])
    logger.info("Logged feedback: %s for query '%s'", rating, clean_q[:30])

